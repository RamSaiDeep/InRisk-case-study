"""Exports: the workbook and the report.

No network. The PDF test needs a local headless browser and is skipped where
there is none.
"""

from __future__ import annotations

from io import BytesIO

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from api.contract_terms import DEFAULT_TERMS
from api.main import app
from api.report import BROWSERS, build_markdown
from api.workbook import _payout_formula, _premium_layout, build_workbook
from data_layer import cache
from pricing_engine import REFERENCE_INPUTS, price

from test_api import REFERENCE_PAYLOAD

HAS_BROWSER = any(path.exists() for path in BROWSERS)


@pytest.fixture
def client(tmp_path, monkeypatch):
    series = tmp_path / "series"
    series.mkdir()
    monkeypatch.setattr(cache, "CACHE_DIR", tmp_path)
    monkeypatch.setattr(cache, "METADATA_FILE", tmp_path / "metadata.json")
    monkeypatch.setattr(cache, "SERIES_DIR", series)
    return TestClient(app)


@pytest.fixture
def frame(reference_daily):
    return pd.DataFrame(
        {
            "date": pd.to_datetime([r.day for r in reference_daily]),
            "latitude": 22.999693294401865,
            "longitude": 72.601155692309760,
            "ssrd_j_m2": [r.ssrd * 3_600_000 for r in reference_daily],
            "ssrd_kwh_m2": [r.ssrd for r in reference_daily],
        }
    )


@pytest.fixture
def seeded(client, frame):
    cache.write_cache(
        "380006",
        frame,
        location={
            "pixel_lat": 22.999693294401865,
            "pixel_lon": 72.601155692309760,
            "attributes": {"Office_Name": "Ellisbridge SO"},
        },
        start_year=2005,
        end_year=2024,
    )
    return client


@pytest.fixture
def result(reference_daily):
    return price(reference_daily, REFERENCE_INPUTS)


META = {
    "pincode": "380006",
    "start_year": 2005,
    "end_year": 2024,
    "pixel_lat": 22.999693294401865,
    "pixel_lon": 72.601155692309760,
    "attributes": {"Office_Name": "Ellisbridge SO"},
}


# --- the payout formula: the one piece of real logic in the workbook -----


def test_payout_formula_is_nested_if_not_ifs():
    """IFS is a future function: written plainly it has to be stored as
    _xlfn.IFS, and reads as #NAME? otherwise. Nested IF needs no prefix."""
    formula = _payout_formula(6, 3, _premium_layout(3), "Inputs!$C$14")
    assert "IFS(" not in formula
    assert formula.count("IF(") == 3


def test_payout_formula_tests_the_trigger_first_then_each_boundary():
    formula = _payout_formula(6, 3, _premium_layout(3), "Inputs!$C$14")
    # Outermost test is the trigger; a year above it falls through to 0.
    assert formula.startswith("=IF(E6<Inputs!$C$14,")
    assert formula.endswith(",0)")
    # Most severe payout sits innermost, behind both boundary tests.
    assert "IF(E6<Premium!$C$9,Premium!$C$12,Premium!$C$11)" in formula


def test_payout_formula_generalises_to_one_level():
    formula = _payout_formula(6, 1, _premium_layout(1), "Inputs!$C$12")
    assert formula == "=IF(E6<Inputs!$C$12,Premium!$C$8,0)"


# --- workbook structure --------------------------------------------------


def test_workbook_has_the_six_sheets_in_order(result, frame):
    book = openpyxl.load_workbook(build_workbook(result, REFERENCE_INPUTS, frame, META, DEFAULT_TERMS))
    assert book.sheetnames == [
        "README",
        "Inputs",
        "Raw Data",
        "Index & Backtest",
        "Premium",
        "Assumptions & Limitations",
    ]


def test_workbook_carries_every_raw_reading(result, frame):
    book = openpyxl.load_workbook(build_workbook(result, REFERENCE_INPUTS, frame, META, DEFAULT_TERMS))
    sheet = book["Raw Data"]
    assert sheet.max_row - 1 == 7305
    assert [cell.value for cell in sheet[1]] == [
        "date",
        "latitude",
        "longitude",
        "ssrd_j_m2",
        "ssrd_kwh_m2",
    ]


def test_workbook_is_formula_driven_not_baked(result, frame):
    """Changing Inputs must move everything downstream, like the reference
    workbook - so no sheet may carry a hard-coded result."""
    book = openpyxl.load_workbook(build_workbook(result, REFERENCE_INPUTS, frame, META, DEFAULT_TERMS))
    premium = book["Premium"]
    layout = _premium_layout(3)
    for key in ("sigma", "boundary_0", "payout_0", "burn_cost", "gross"):
        value = premium.cell(row=layout[key], column=3).value
        assert isinstance(value, str) and value.startswith("="), key
    backtest = book["Index & Backtest"]
    assert str(backtest["E6"].value).startswith("=D6*Inputs!")
    assert "STDEV" in str(premium.cell(row=layout["payout_stdev"], column=3).value)


def test_workbook_gross_up_divides(result, frame):
    book = openpyxl.load_workbook(build_workbook(result, REFERENCE_INPUTS, frame, META, DEFAULT_TERMS))
    layout = _premium_layout(3)
    formula = book["Premium"].cell(row=layout["gross"], column=3).value
    assert "/(1-Inputs!" in formula


def test_workbook_records_where_the_data_came_from(result, frame):
    book = openpyxl.load_workbook(build_workbook(result, REFERENCE_INPUTS, frame, META, DEFAULT_TERMS))
    readme = "\n".join(
        str(cell.value) for row in book["README"].iter_rows() for cell in row if cell.value
    )
    assert "380006" in readme
    assert "22.999693" in readme
    assert "2005-2024" in readme


def test_workbook_names_the_bands_that_never_paid(result, frame):
    book = openpyxl.load_workbook(build_workbook(result, REFERENCE_INPUTS, frame, META, DEFAULT_TERMS))
    text = "\n".join(
        str(cell.value)
        for row in book["Assumptions & Limitations"].iter_rows()
        for cell in row
        if cell.value
    )
    assert "Severe" in text  # no year on record reached it
    assert "2024" in text  # the closest call


# --- report --------------------------------------------------------------


def test_report_states_the_schedule_and_the_numbers(result):
    text = build_markdown(result, REFERENCE_INPUTS, META, DEFAULT_TERMS)
    assert "₹137.48" in text
    assert "₹608.43" in text
    assert "4,515.35" in text  # the exit
    assert "Calculation Agent" in text  # wording, not a number
    assert "2011 (Mild)" in text and "2019 (Moderate)" in text
    assert "+1.24%" in text  # mean generation against the trigger
    assert "2024" in text and "1.71 kWh from an edge" in text


def test_report_tables_are_single_blocks(result):
    """A pipe table split across blank lines renders as paragraphs of pipes."""
    text = build_markdown(result, REFERENCE_INPUTS, META, DEFAULT_TERMS)
    for block in text.split("\n\n"):
        if block.startswith("|"):
            assert block.count("\n") >= 2, block[:60]


def test_report_says_when_a_band_was_never_observed(result):
    text = build_markdown(result, REFERENCE_INPUTS, META, DEFAULT_TERMS)
    assert "No year on record reached the Severe band" in text


# --- through the API -----------------------------------------------------


def test_xlsx_endpoint_returns_a_workbook(seeded):
    response = seeded.post("/api/export/xlsx", json=REFERENCE_PAYLOAD)
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "Solar_Parametric_380006_2005-2024.xlsx" in response.headers["content-disposition"]
    book = openpyxl.load_workbook(BytesIO(response.content))
    assert book["Raw Data"].max_row - 1 == 7305


def test_xlsx_export_honours_the_year_range(seeded):
    payload = {**REFERENCE_PAYLOAD, "start_year": 2019, "end_year": 2024}
    response = seeded.post("/api/export/xlsx", json=payload)
    book = openpyxl.load_workbook(BytesIO(response.content))
    assert book["Raw Data"].max_row - 1 == 2192  # 2019-2024, with two leap years
    assert book["Index & Backtest"]["A6"].value == 2019


def test_exports_refuse_an_unfetched_pincode(seeded):
    for route in ("/api/export/xlsx", "/api/export/report"):
        response = seeded.post(route, json=REFERENCE_PAYLOAD | {"pincode": "110001"})
        assert response.status_code == 409


def test_exports_reject_bad_inputs_the_same_way_pricing_does(seeded):
    payload = {
        **REFERENCE_PAYLOAD,
        "inputs": {**REFERENCE_PAYLOAD["inputs"], "boundary_sigmas": [1.0, 0.5]},
    }
    assert seeded.post("/api/export/xlsx", json=payload).status_code == 422


@pytest.mark.skipif(not HAS_BROWSER, reason="needs a local headless Chrome or Edge")
def test_report_endpoint_returns_a_pdf(seeded):
    response = seeded.post("/api/export/report", json=REFERENCE_PAYLOAD)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert len(response.content) > 20_000
